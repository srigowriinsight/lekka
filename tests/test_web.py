"""Every page loads, and the main popups work end to end through HTTP."""
from datetime import date
from decimal import Decimal

import pytest

from app.demo import load_demo
from app.extensions import db
from app.models import CashflowEntry, Head, RecurringTemplate
from app.services.balances import closing_on

F = {"X-Requested-With": "fetch"}


@pytest.fixture
def demo_client(client):
    load_demo(shift=False)
    db.session.commit()
    return client


def by_title(t):
    return CashflowEntry.query.filter_by(title=t).first()


def test_login_required(app):
    c = app.test_client()
    r = c.get("/")
    assert r.status_code == 302 and "/login" in r.headers["Location"]


@pytest.mark.parametrize("path", [
    "/", "/?freq=week", "/?freq=month", "/?freq=quarter", "/?freq=year", "/?freq=day&hide_empty=0",
    "/?freq=month&compact=1&expand=1", "/?freq=day&filter=pending", "/?freq=day&filter=settled",
    "/expenses", "/revenue", "/expenses?status=overdue", "/revenue?status=overdue", "/expenses?q=PO-V",
    "/review", "/templates/", "/master/", "/master/?tab=cost_centres", "/master/?tab=counterparties",
    "/account", "/history", "/adjustments?a=2026-11-01&b=2026-11-25", "/templates/new", "/entries/new?direction=IN",
    "/master/heads/new",
])
def test_pages_load(demo_client, path):
    r = demo_client.get(path)
    assert r.status_code == 200, path


def test_grid_shows_worked_example(demo_client):
    html = demo_client.get("/?freq=month&start=2026-11-01&end=2027-01-31").get_data(as_text=True)
    for figure in ("20,65,000", "11,40,000", "1,65,000", "8,40,000"):
        assert figure in html


def test_popups_for_an_item(demo_client):
    i1 = by_title("Instalment 1")
    for path in (f"/entries/{i1.id}", f"/entries/{i1.id}/edit", f"/entries/{i1.id}/split",
                 f"/entries/{i1.id}/settle", f"/entries/{i1.id}/reschedule", f"/entries/{i1.id}/cancel"):
        assert demo_client.get(path, headers=F).status_code == 200, path
    po = by_title("PO-V-221 components")
    assert demo_client.get(f"/entries/{po.id}/split", headers=F).status_code == 200
    head = i1.head_id
    r = demo_client.get(f"/grid/cell?head={head}&dir=OUT&a=2026-11-01&b=2026-11-30", headers=F)
    assert "Instalment 1" in r.get_data(as_text=True)


def test_create_and_settle_via_http(demo_client):
    head = Head.query.filter_by(name="Scrap").first()
    r = demo_client.post("/entries/new", headers=F, data={
        "direction": "IN", "head_id": head.id, "title": "Scrap sale", "amount": "12,000",
        "planned_date": "2026-12-03", "remarks": "copper scrap", "next": "/revenue"})
    assert r.get_json() == {"redirect": "/revenue"}
    e = by_title("Scrap sale")
    assert e.amount == Decimal("12000") and e.remarks == "copper scrap"
    r = demo_client.post(f"/entries/{e.id}/settle", headers=F, data={"when": "today", "version": e.version})
    assert "redirect" in r.get_json()
    db.session.refresh(e)
    assert e.is_settled and e.settled_on == date(2026, 11, 25)


def test_errors_come_back_to_the_popup(demo_client):
    r = demo_client.post("/entries/new", headers=F, data={"direction": "OUT", "head_id": "", "title": "x",
                                                         "amount": "100", "planned_date": "2026-12-01"})
    assert r.status_code == 400 and "head" in r.get_json()["error"].lower()


def test_new_head_inline(demo_client):
    demo_client.post("/entries/new", headers=F, data={"direction": "OUT", "new_head": "Packing material",
                                                     "title": "Boxes", "amount": "5000", "planned_date": "2026-12-01"})
    assert Head.query.filter_by(name="Packing material").first() is not None


def test_link_prompt_then_link(demo_client):
    inv = by_title("Invoice INV-388")
    data = {"direction": "IN", "head_id": inv.head_id, "title": "Payment from customer", "amount": "2,00,000",
            "planned_date": "2026-11-25", "already_settled": "1"}
    r = demo_client.post("/entries/new", headers=F, data=data)
    assert "Invoice INV-388" in r.get_json()["html"]
    r = demo_client.post("/entries/new", headers=F, data={**data, "link_choice": str(inv.id)})
    assert "redirect" in r.get_json()
    db.session.refresh(inv)
    assert inv.is_settled and inv.settled_on == date(2026, 11, 25)
    assert CashflowEntry.query.filter_by(title="Payment from customer").first() is None


def test_split_via_http(demo_client):
    i2 = by_title("Instalment 2")
    r = demo_client.post(f"/entries/{i2.id}/split", headers=F, data={
        "part_id": ["", ""], "part_title": ["Instalment 2a", "Instalment 2b"],
        "part_amount": ["1,50,000", "1,00,000"], "part_date": ["2026-12-10", "2026-12-20"],
        "part_remarks": ["", ""], "version": i2.version})
    assert "redirect" in r.get_json()
    db.session.refresh(i2)
    assert [c.title for c in i2.children] == ["Instalment 2a", "Instalment 2b"]


def test_cancel_needs_choice_via_http(demo_client):
    i2 = by_title("Instalment 2")
    r = demo_client.post(f"/entries/{i2.id}/cancel", headers=F, data={"reason": "Order reduced"})
    assert r.status_code == 400
    r = demo_client.post(f"/entries/{i2.id}/cancel", headers=F, data={"reason": "Order reduced", "mode": "shrink"})
    assert "redirect" in r.get_json()


def test_adjustment_via_http(demo_client):
    r = demo_client.post("/adjustments/new", headers=F, data={"adj_date": "2026-11-25", "sign": "-",
                                                              "amount": "350", "remarks": "Bank charges"})
    assert "redirect" in r.get_json()
    assert closing_on(date(2026, 11, 25)) == Decimal("2964650")


def test_template_via_http(demo_client):
    head = Head.query.filter_by(name="Factory rent").first()
    r = demo_client.post("/templates/new", headers=F, data={
        "title": "Warehouse rent", "head_id": head.id, "direction": "OUT", "amount": "40,000",
        "freq": "MONTHLY", "interval_n": "1", "anchor_day": "-1", "start_date": "2026-12-01"})
    assert "redirect" in r.get_json()
    t = RecurringTemplate.query.filter_by(title="Warehouse rent").one()
    assert demo_client.get(f"/templates/{t.id}/edit", headers=F).status_code == 200


def test_export_excel(demo_client):
    r = demo_client.get("/grid/export?freq=month&start=2026-11-01&end=2027-01-31")
    assert r.status_code == 200 and r.data[:2] == b"PK"
    from io import BytesIO
    from openpyxl import load_workbook
    ws = load_workbook(BytesIO(r.data)).active
    closing_row = [row for row in ws.iter_rows(values_only=True) if row[0] == "Closing balance"][0]
    assert closing_row[1:4] == (2065000.0, 1140000.0, 165000.0)


def test_unsafe_next_ignored(demo_client):
    i1 = by_title("Instalment 1")
    r = demo_client.post(f"/entries/{i1.id}/reschedule", data={"new_date": "2026-12-02",
                                                              "next": "https://evil.example/"})
    assert r.status_code == 302 and "evil" not in r.headers["Location"]


def test_day_view_shows_every_day(demo_client):
    # Default Day view: the last 10 days, today (25 Nov in tests) and the next 5, empty days included.
    html = demo_client.get("/").get_data(as_text=True)
    for label in ("Sun 15 Nov", "Mon 16 Nov", "Wed 25 Nov", "Mon 30 Nov"):
        assert label in html, label
    assert "Tue 1 Dec" not in html and "Sat 14 Nov" not in html
    assert "Hide empty days" in html
    # A custom range shows every day in it, and Custom is marked as the chosen view.
    html = demo_client.get("/?freq=day&custom=1&start=2026-11-02&end=2026-11-06").get_data(as_text=True)
    for label in ("Mon 2 Nov", "Tue 3 Nov", "Wed 4 Nov", "Thu 5 Nov", "Fri 6 Nov"):
        assert label in html, label
    assert 'class="on" title="Day columns' in html


def test_dates_before_opening_are_explained(client):
    # A fresh account opens today (25 Nov in tests) until it is set up; the grid says so.
    from app.models import AuditLog
    AuditLog.query.filter_by(table_name="account").delete()
    db.session.commit()
    assert "Set up the account first" in client.get("/").get_data(as_text=True)
    h = Head(name="Salaries")
    db.session.add(h)
    db.session.commit()
    # Account opened on 25 Nov: an item planned on 23 Nov is refused with the reason.
    client.post("/account", data={"name": "Main", "opening_balance": "10,00,000", "opening_date": "2026-11-25"})
    r = client.post("/entries/new", data={"direction": "OUT", "head_id": h.id, "title": "Contract labour",
                                          "amount": "40000", "planned_date": "2026-11-23"}, headers=F)
    assert "before the account's opening date (25 Nov 2026)" in r.get_data(as_text=True)
    assert CashflowEntry.query.filter_by(title="Contract labour").count() == 0
    assert "Set up the account first" not in client.get("/").get_data(as_text=True)
    # Moving the opening date earlier makes the same item work: shown in the grid and counted once ticked.
    client.post("/account", data={"name": "Main", "opening_balance": "10,00,000", "opening_date": "2026-11-15"})
    client.post("/entries/new", data={"direction": "OUT", "head_id": h.id, "title": "Contract labour",
                                      "amount": "40000", "planned_date": "2026-11-23"}, headers=F)
    e = by_title("Contract labour")
    assert "Contract labour" in client.get("/review").get_data(as_text=True)
    client.post(f"/entries/{e.id}/settle", data={"when": "pick", "settled_on": "2026-11-24"}, headers=F)
    assert closing_on(date(2026, 11, 25)) == Decimal("960000")


def test_quick_amount_edit_from_grid(demo_client):
    # Double-clicking a figure posts just the new amount for that item.
    e = CashflowEntry.query.filter(CashflowEntry.is_settled.is_(False),
                                   CashflowEntry.planned_date > date(2026, 11, 25)).first()
    e = e.leaves()[0]
    html = demo_client.get("/").get_data(as_text=True)
    assert "data-edit=" in html and "Double-click to change the amount" in html
    r = demo_client.post(f"/entries/{e.id}/amount", data={"amount": "1,23,456", "version": e.version}, headers=F)
    assert r.get_json().get("redirect")
    assert by_title(e.title).amount == Decimal("123456")
    # A stale version (changed elsewhere) is refused with a clear message.
    r = demo_client.post(f"/entries/{e.id}/amount", data={"amount": "5000", "version": 0}, headers=F)
    assert "changed elsewhere" in r.get_json()["error"]


def test_recurring_template_more_and_delete(demo_client):
    h = Head.query.filter_by(name="Salaries").first()
    demo_client.post("/templates/new", data={"title": "Staff salary", "head_id": h.id, "direction": "OUT",
                                             "amount": "3,00,000", "freq": "MONTHLY", "interval_n": "1",
                                             "start_date": "2026-11-30"}, headers=F)
    t = RecurringTemplate.query.filter_by(title="Staff salary").first()
    html = demo_client.get(f"/templates/{t.id}/more", headers=F).get_data(as_text=True)
    assert "End series" in html and "Delete recurring template" in html
    r = demo_client.post(f"/templates/{t.id}/action", data={"do": "delete"}, headers=F)
    assert r.get_json().get("redirect")
    assert db.session.get(RecurringTemplate, t.id) is None
    page = demo_client.get("/templates/").get_data(as_text=True)
    assert "Recurring templates" in page and "+ New recurring template" in page


def test_cost_centres_hidden(demo_client):
    assert "Cost centres" not in demo_client.get("/master/").get_data(as_text=True)
    assert "All cost centres" not in demo_client.get("/").get_data(as_text=True)
    assert "Counterparty" not in demo_client.get("/entries/new?direction=IN", headers=F).get_data(as_text=True)


def test_history_is_compact_for_recurring_templates(demo_client):
    from app.models import AuditLog
    h = Head.query.filter_by(name="Salaries").first()
    before = AuditLog.query.count()
    demo_client.post("/templates/new", data={"title": "subscription", "head_id": h.id, "direction": "OUT",
                                             "amount": "1,000", "freq": "MONTHLY", "interval_n": "1",
                                             "start_date": "2026-11-26"}, headers=F)
    # One log line for the template and all its dates, not one per date.
    assert AuditLog.query.count() - before == 1
    page = demo_client.get("/history").get_data(as_text=True)
    assert "Recurring template added" in page and "12 dates planned: 26 Nov 2026 – 26 Oct 2027" in page
    assert "UTC" not in page


def test_history_folds_old_per_date_lines(app):
    # Data written by the earlier version (one 'generate' line per date) still shows as one line.
    from datetime import datetime
    from app.models import AuditLog
    from app.services.history import build
    at = datetime(2026, 10, 24, 5, 41)
    rows = [AuditLog(table_name="cashflow_entry", record_id=70 + i, action="generate", at=at,
                     new_json={"title": "subscription", "planned_date": f"2027-0{i}-24"}) for i in range(5, 0, -1)]
    rows.append(AuditLog(table_name="recurring_template", record_id=3, action="create", at=at,
                         new_json={"title": "subscription"}))
    out = build(rows)
    assert len(out) == 1
    assert out[0].what == "Recurring template added" and out[0].remarks == "5 dates planned: 24 Jan 2027 – 24 May 2027"


def test_history_filters(demo_client):
    from app.models import AuditLog
    page = demo_client.get("/history").get_data(as_text=True)
    assert "Search name or details" in page
    e = CashflowEntry.query.filter(CashflowEntry.is_settled.is_(False), CashflowEntry.planned_date.isnot(None),
                                   CashflowEntry.planned_date <= date(2026, 11, 25)).first()
    demo_client.post(f"/entries/{e.id}/settle", data={"when": "today"}, headers=F)
    ticked = demo_client.get("/history?what=ticked").get_data(as_text=True)
    assert "<td>Ticked</td>" in ticked and "<td>Added</td>" not in ticked
    items_only = demo_client.get("/history?kind=templates").get_data(as_text=True)
    assert "Nothing matches these filters." in items_only            # the demo has no recurring templates
    found = demo_client.get("/history?q=PO-V-221").get_data(as_text=True)
    assert "PO-V-221 components" in found and "Calibration contract" not in found
    assert demo_client.get("/history?from=2099-01-01").status_code == 200
    assert demo_client.get("/history?from=bad&page=x").status_code == 200


def test_enter_saves_new_item(demo_client):
    # The plain Save button comes first in the form, so pressing Enter saves instead of opening Split.
    html = demo_client.get("/entries/new?direction=OUT", headers=F).get_data(as_text=True)
    assert html.index('class="primary" style="order:2">Save<') < html.index('name="then_split"')


def test_delete_hides_item_and_cancel_keeps_it_visible(demo_client):
    h = Head.query.filter_by(name="Salaries").first()
    for title in ("Wrong entry", "Bonus that won't happen"):
        demo_client.post("/entries/new", data={"direction": "OUT", "head_id": h.id, "title": title,
                                               "amount": "10,000", "planned_date": "2026-11-27"}, headers=F)
    before = closing_on(date(2026, 11, 30))
    wrong, bonus = by_title("Wrong entry"), by_title("Bonus that won't happen")
    # Delete: gone from the grid, the lists and the balances.
    r = demo_client.post(f"/entries/{wrong.id}/delete", data={"version": wrong.version}, headers=F)
    assert r.get_json().get("redirect")
    assert closing_on(date(2026, 11, 30)) == before + 10000
    demo_client.get("/")                       # show (and clear) the "'Wrong entry' deleted." message
    grid = demo_client.get("/?freq=day&start=2026-11-20&end=2026-11-30&expand=1").get_data(as_text=True)
    expense = demo_client.get("/expenses?start=2026-11-01&end=2026-11-30").get_data(as_text=True)
    assert "Wrong entry" not in grid and "Wrong entry" not in expense
    assert demo_client.post(f"/entries/{wrong.id}/settle", data={"when": "today"}, headers=F).status_code == 400
    assert "Deleted" in demo_client.get("/history").get_data(as_text=True)
    # Cancel: still shown, struck through, and not counted.
    demo_client.post(f"/entries/{bonus.id}/cancel", data={"reason": "Not approved"}, headers=F)
    assert closing_on(date(2026, 11, 30)) == before + 20000
    grid = demo_client.get("/?freq=day&start=2026-11-20&end=2026-11-30&expand=1").get_data(as_text=True)
    expense = demo_client.get("/expenses?start=2026-11-01&end=2026-11-30").get_data(as_text=True)
    assert "Bonus that won" in grid and "cancelled-amt" in grid and "c-CANCELLED" in grid
    assert "Bonus that won" in expense


def test_deleted_recurring_date_is_not_recreated(app, head):
    from app.services.entries import delete_entry
    from app.services.recurrence import create_template, extend_all, occurrences
    t = create_template(dict(title="Rent", head_id=head("Rent").id, direction="OUT", amount="1000",
                             freq="MONTHLY", interval_n=1, start_date="2026-12-01"))
    first = occurrences(t)[0]
    n = len(occurrences(t))
    delete_entry(first)
    extend_all()
    assert len(occurrences(t)) == n - 1 and first.occurrence_date not in [o.occurrence_date for o in occurrences(t)]


def test_split_part_offers_delete_whole_item(demo_client):
    part = by_title("Instalment 2")                    # a part two levels down in the demo
    html = demo_client.get(f"/grid/cell?head={part.head_id}&dir={part.direction}"
                           f"&a={part.planned_date}&b={part.planned_date}", headers=F).get_data(as_text=True)
    top = part.ancestors()[-1]
    assert "Delete whole item" in html and f"/entries/{top.id}/delete" in html


def test_correct_amount_of_ticked_item(demo_client):
    # Double-click a settled figure: untick, change, tick again on the same date, all in one step.
    e = CashflowEntry.query.filter(CashflowEntry.is_settled.is_(True), CashflowEntry.parent_id.is_(None)).first()
    paid_on, old_amount = e.settled_on, e.amount
    before = closing_on(paid_on)
    html = demo_client.get(f"/?freq=day&start={paid_on}&end={paid_on}").get_data(as_text=True)
    assert "(it stays ticked)" in html
    r = demo_client.post(f"/entries/{e.id}/amount", data={"amount": str(old_amount + 1000), "version": e.version},
                         headers=F)
    assert r.get_json().get("redirect")
    e = db.session.get(CashflowEntry, e.id)
    assert e.is_settled and e.settled_on == paid_on and e.amount == old_amount + 1000
    sign = 1 if e.direction == "IN" else -1
    assert closing_on(paid_on) == before + sign * 1000
    assert "Amount corrected from" in demo_client.get("/history").get_data(as_text=True)


def test_past_entry_can_tick_the_future_plan(demo_client):
    # A plan for 30 Nov exists. The money actually went on 24 Nov and was entered as a new item, then ticked.
    h = Head.query.filter_by(name="Salaries").first()
    demo_client.post("/entries/new", data={"direction": "OUT", "head_id": h.id, "title": "Bonus",
                                           "amount": "50,000", "planned_date": "2026-11-30"}, headers=F)
    before = closing_on(date(2026, 11, 30))
    demo_client.post("/entries/new", data={"direction": "OUT", "head_id": h.id, "title": "Bonus paid",
                                           "amount": "50,000", "planned_date": "2026-11-24"}, headers=F)
    paid, plan = by_title("Bonus paid"), by_title("Bonus")
    popup = demo_client.get(f"/entries/{paid.id}/settle", headers=F).get_data(as_text=True)
    assert "Is this the money for a plan you already have?" in popup and f'value="{plan.id}"' in popup
    r = demo_client.post(f"/entries/{paid.id}/settle", data={"when": "planned", "link_to": plan.id,
                                                              "version": paid.version}, headers=F)
    assert r.get_json().get("redirect")
    plan, paid = by_title("Bonus"), by_title("Bonus paid")
    assert plan.is_settled and plan.settled_on == date(2026, 11, 24) and paid.is_deleted
    assert closing_on(date(2026, 11, 30)) == before          # counted once, not twice
    assert closing_on(date(2026, 11, 24)) < closing_on(date(2026, 11, 23))


def test_new_item_in_the_past_is_marked_already_paid(demo_client):
    html = demo_client.get("/entries/new?direction=OUT&date=2026-11-20", headers=F).get_data(as_text=True)
    assert 'name="already_settled" value="1" checked' in html
    html = demo_client.get("/entries/new?direction=OUT&date=2026-11-28", headers=F).get_data(as_text=True)
    assert 'name="already_settled" value="1" checked' not in html


def test_expanded_item_line_opens_the_cell_popup(demo_client):
    html = demo_client.get("/?freq=day&expand=1").get_data(as_text=True)
    assert "/grid/cell?" in html and "entry=" in html
    e = [x for x in CashflowEntry.query.filter(CashflowEntry.is_settled.is_(False), CashflowEntry.parent_id.is_(None),
                                                CashflowEntry.planned_date.isnot(None)).all() if x.is_leaf][0]
    pop = demo_client.get(f"/grid/cell?head={e.head_id}&dir={e.direction}&a={e.planned_date}&b={e.planned_date}"
                          f"&entry={e.id}", headers=F).get_data(as_text=True)
    assert e.title in pop and ">Tick<" in pop and ">Reschedule<" in pop and ">Delete<" in pop


def test_reset_a_period_removes_only_what_was_entered_then(demo_client):
    import time
    from datetime import datetime, timedelta, timezone
    from app.models import BalanceAdjustment
    ist = lambda d: (d.replace(tzinfo=timezone.utc) + timedelta(hours=5, minutes=30)).strftime("%Y-%m-%dT%H:%M")
    old = by_title("Invoice INV-388")
    before = closing_on(date(2026, 11, 30))
    time.sleep(0.01)
    start = datetime.utcnow().replace(second=0, microsecond=0) + timedelta(minutes=1)   # after the demo data
    h = Head.query.filter_by(name="Salaries").first()
    # Pretend these were entered inside the window by stamping their entry time.
    demo_client.post("/entries/new", data={"direction": "OUT", "head_id": h.id, "title": "Test expense",
                                           "amount": "5,000", "planned_date": "2026-11-27"}, headers=F)
    demo_client.post("/templates/new", data={"title": "Test rent", "head_id": h.id, "direction": "OUT",
                                             "amount": "1,000", "freq": "MONTHLY", "interval_n": "1",
                                             "start_date": "2026-11-28"}, headers=F)
    demo_client.post("/adjustments/new", data={"adj_date": "2026-11-25", "sign": "-", "amount": "100"}, headers=F)
    for obj in [by_title("Test expense"), RecurringTemplate.query.filter_by(title="Test rent").first(),
                BalanceAdjustment.query.order_by(BalanceAdjustment.id.desc()).first()]:
        obj.created_at = start + timedelta(minutes=5)
    old_title = old.title
    edit_amount = old.amount
    db.session.commit()
    window = {"start": ist(start), "end": ist(start + timedelta(minutes=30))}
    pop = demo_client.post("/history/reset", data={**window, "step": "preview"}, headers=F).get_json()["html"]
    assert "Test expense" in pop and "Test rent" in pop and "Adjustment on" in pop and old_title not in pop
    r = demo_client.post("/history/reset", data={**window, "step": "run"}, headers=F)
    assert "Tick the box" in r.get_json()["error"]
    r = demo_client.post("/history/reset", data={**window, "step": "run", "sure": "1"}, headers=F)
    assert r.get_json().get("redirect")
    assert by_title("Test expense").is_deleted
    assert RecurringTemplate.query.filter_by(title="Test rent").first() is None
    assert not CashflowEntry.query.filter(CashflowEntry.title == "Test rent",
                                          CashflowEntry.lifecycle != "DELETED").count()
    assert not by_title(old_title).is_deleted and by_title(old_title).amount == edit_amount
    assert closing_on(date(2026, 11, 30)) == before
    assert "Period reset" in demo_client.get("/history").get_data(as_text=True)
