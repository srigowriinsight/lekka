"""The grid page (one page for viewing), its cell drawer, the overdue review and Excel export."""
from datetime import date, timedelta
from io import BytesIO

from dateutil.relativedelta import relativedelta
from flask import Blueprint, render_template, request, send_file
from flask_login import login_required

from ..extensions import db
from ..models import IN, OUT, CostCentre, Head
from ..services.balances import FREQS, build_grid, cell_items, due_today_items, fy_start, overdue_items
from ..services.common import LekkaError, inr, parse_date, today
from ..services.recurrence import extend_all
from ._util import fragment

bp = Blueprint("grid", __name__)


@bp.app_context_processor
def nav_context():
    from flask_login import current_user
    if not getattr(current_user, "is_authenticated", False):
        return {}
    try:
        return {"overdue_count": len(overdue_items())}
    except Exception:
        return {"overdue_count": 0}


def default_range(freq, t):
    if freq == "day":
        # Like the Screens mockup (15–30 Nov with today 25 Nov): the last 10 days, today, and the next 5.
        return t - timedelta(days=10), t + timedelta(days=5)
    if freq == "week":
        start = t - timedelta(days=t.weekday()) - timedelta(weeks=2)
        return start, start + timedelta(weeks=12) - timedelta(days=1)
    if freq == "month":
        start = t.replace(day=1) - relativedelta(months=1)
        return start, start + relativedelta(months=6) - timedelta(days=1)
    if freq == "quarter":
        return fy_start(t), fy_start(t) + relativedelta(years=1) - timedelta(days=1)
    if freq == "year":
        return fy_start(t) - relativedelta(years=1), fy_start(t) + relativedelta(years=2) - timedelta(days=1)
    return t.replace(day=1), t.replace(day=1) + relativedelta(months=1) - timedelta(days=1)


def grid_params():
    t = today()
    freq = request.args.get("freq", "day")
    if freq not in FREQS:
        freq = "day"
    a0, b0 = default_range(freq, t)
    try:
        a = parse_date(request.args["start"]) if request.args.get("start") else a0
        b = parse_date(request.args["end"]) if request.args.get("end") else b0
    except LekkaError:
        a, b = a0, b0
    if b < a:
        a, b = b, a
    if freq == "day" and (b - a).days > 400:
        b = a + timedelta(days=400)
    return dict(
        freq=freq, start=a, end=b,
        filter=request.args.get("filter", "all"),
        # Every day in the range gets a column unless the user asks to hide empty ones.
        hide_empty=request.args.get("hide_empty", "0") == "1",
        custom=freq == "day" and request.args.get("custom") == "1",
        expand_all=request.args.get("expand") == "1",
        compact=request.args.get("compact") == "1",
        cost_centre_id=request.args.get("cc") or None,
    )


def _shift(p, direction):
    span = (p["end"] - p["start"]) + timedelta(days=1)
    if p["freq"] in ("month", "quarter", "year"):
        months = {"month": 6, "quarter": 12, "year": 36}[p["freq"]]
        delta = relativedelta(months=months)
        return (p["start"] + delta * direction, p["end"] + delta * direction)
    return (p["start"] + span * direction, p["end"] + span * direction)


@bp.route("/")
@login_required
def index():
    if extend_all():
        db.session.commit()
    p = grid_params()
    g = build_grid(p["start"], p["end"], p["freq"], item_filter=p["filter"],
                   hide_empty=p["hide_empty"], cost_centre_id=p["cost_centre_id"])
    prev_range, next_range = _shift(p, -1), _shift(p, 1)
    from ..models import AuditLog
    from ..services.entries import get_account
    acc = get_account()
    account_set = AuditLog.query.filter_by(table_name="account").first() is not None
    return render_template("grid.html", g=g, p=p, IN=IN, OUT=OUT, prev_range=prev_range, acc=acc,
                           account_set=account_set,
                           next_range=next_range, cost_centres=CostCentre.query.filter_by(is_active=True).all(),
                           heads=Head.query.filter_by(is_active=True).order_by(Head.sort_order, Head.name).all())


@bp.route("/grid/cell")
@login_required
def cell():
    head = db.session.get(Head, int(request.args["head"]))
    direction = request.args["dir"]
    a, b = parse_date(request.args["a"]), parse_date(request.args["b"])
    if request.args.get("entry"):        # an item's own line (Expand all): the same popup, for that item only
        from ..services.entries import get_entry
        items = [get_entry(request.args["entry"])]
    else:
        items = cell_items(head.id, direction, a, b)
    return fragment("partials/cell.html", head=head, direction=direction, a=a, b=b, items=items)


@bp.route("/review")
@login_required
def review():
    from ..services.entries import get_account
    return render_template("review.html", overdue=overdue_items(), due_today=due_today_items(),
                           opening_date=get_account().opening_date)


@bp.route("/grid/export")
@login_required
def export():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    p = grid_params()
    g = build_grid(p["start"], p["end"], p["freq"], item_filter=p["filter"],
                   hide_empty=p["hide_empty"], cost_centre_id=p["cost_centre_id"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Lekka"
    bold = Font(bold=True)
    band = PatternFill("solid", fgColor="EEEEEE")
    ws.append(["Lekka cashflow", f"{p['start']:%d %b %Y} – {p['end']:%d %b %Y}", f"View: {p['freq']}",
               f"Exported {today():%d %b %Y}"])
    ws.append([])
    ws.append([""] + [c.label for c in g.columns])
    ws.append(["Basis"] + [c.basis for c in g.columns])
    ws.append(["Opening balance"] + [float(c.opening) for c in g.columns])

    def block(name, direction, total_attr):
        ws.append([name])
        ws.cell(ws.max_row, 1).fill = band
        for head, cells, items in g.rows[direction]:
            ws.append([head.name] + [float(cells[i].amount) if i in cells else None for i in range(len(g.columns))])
            if p["expand_all"]:
                for it in items:
                    e = it["entry"]
                    ws.append(["    " + e.title] + [float(it["cells"][i]["amount"]) if i in it["cells"] else None
                                                    for i in range(len(g.columns))])
        ws.append([f"Total {name.lower()}"] + [float(getattr(c, total_attr)) for c in g.columns])
        ws.cell(ws.max_row, 1).font = bold

    block("Inflow", IN, "total_in")
    block("Outflow", OUT, "total_out")
    ws.append(["Adjustment"] + [float(c.adjustment) for c in g.columns])
    ws.append(["Closing balance"] + [float(c.closing) for c in g.columns])
    for cell in ws[ws.max_row]:
        cell.font = bold
    for row in ws.iter_rows(min_row=5, min_col=2):
        for c in row:
            c.number_format = r"[>=10000000]##\,##\,##\,##0.00;[>=100000]##\,##\,##0.00;##,##0.00"
    ws.column_dimensions["A"].width = 34
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"lekka_{p['freq']}_{p['start']:%Y%m%d}_{p['end']:%Y%m%d}.xlsx"
    return send_file(buf, as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
